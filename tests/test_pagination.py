"""Every list pages, by the length an admin chose: a big client's screens stay quick."""

import re
from pathlib import Path

import pytest
from django.db.models import Count
from django.urls import reverse
from django.utils import timezone

from apps.adjustments.models import AdjustmentBatch
from apps.core.tenancy import tenant_context
from apps.iam.models import Membership
from apps.org.models import TenantProfile
from apps.org.timezones import tenant_zone_name, zone
from apps.platform.console import sample
from apps.platform.console.web.forms import PlatformSettingsForm
from apps.platform.tenants.models import PlatformSettings
from conftest import login

ROOT = Path(__file__).resolve().parent.parent
ALLOWED = {"apps/core/paging.py",  # the one place
           "apps/mobile/api/v1/views/monitor.py"}  # the mobile contract


def test_no_list_chooses_its_own_page_length():
    offenders = []
    for path in (ROOT / "apps").rglob("*.py"):
        name = path.relative_to(ROOT).as_posix()
        if name not in ALLOWED and "/tests/" not in name and re.search(
                r"\bPaginator\(", path.read_text(encoding="utf-8")):
            offenders.append(name)
    assert offenders == []


@pytest.fixture
def company(tenant_a):
    with tenant_context(tenant_a.id), timezone.override(zone(tenant_zone_name(tenant_a.id))):
        sample.seed(tenant_a, "Sample-Pass-2026", people=40, months=1)
        TenantProfile.objects.update(page_size=10)
    return login(tenant_a, language="en")


def _rows(page: str) -> int:
    body = page[page.index("<tbody"):page.rindex("</tbody>")]
    return len(re.findall(r"<tr[\s>]", body))


@pytest.mark.parametrize("name, args", [
    ("employees", ()), ("attendance", ()), ("report", ("employee_list",)),
    ("report", ("daily_attendance",)), ("pay-item-assignments", ()),
])
def test_a_big_list_shows_the_length_the_admin_chose(tenant_a, company, name, args):
    response = company.get(reverse(name, args=args), HTTP_HOST=f"{tenant_a.slug}.localhost")
    page = response.content.decode()
    assert response.status_code == 200
    assert _rows(page) <= 11  # a report may add its heading row
    assert "page=2" in page


def test_the_days_figures_count_everyone_on_every_page(tenant_a, company):
    url = reverse("attendance")
    host = {"HTTP_HOST": f"{tenant_a.slug}.localhost"}
    first, second = (company.get(f"{url}?page={n}", **host).context for n in (1, 2))
    assert first["summary"] == second["summary"]
    assert len(first["rows"]) > 10 and len(first["page"].object_list) == 10
    assert first["page"].object_list != second["page"].object_list


def test_a_batch_totals_everyone_on_every_page(tenant_a, company):
    with tenant_context(tenant_a.id):
        batch = AdjustmentBatch.objects.annotate(people=Count("rows")).order_by("-people").first()
    url = reverse(f"{batch.kind}-batch", args=[batch.pk])
    context = company.get(url, HTTP_HOST=f"{tenant_a.slug}.localhost").context
    assert context["total"] == sum(row.amount for row in context["rows"])
    assert len(context["page"].object_list) <= 10


def test_a_report_keeps_its_filters_and_its_totals_across_pages(tenant_a, company):
    url = reverse("report", args=["employee_list"])
    page = company.get(f"{url}?page=2", HTTP_HOST=f"{tenant_a.slug}.localhost").content.decode()
    assert 'data-pages="' in page
    assert "format=pdf&amp;page" not in page and "page=2&amp;format" not in page


def test_the_export_has_every_row_whatever_the_page(tenant_a, company):
    url = reverse("report", args=["employee_list"])
    whole = company.get(f"{url}?format=xlsx&page=1", HTTP_HOST=f"{tenant_a.slug}.localhost")
    assert whole.status_code == 200
    with tenant_context(tenant_a.id):
        from io import BytesIO

        from openpyxl import load_workbook

        from apps.employees.models import Employee

        sheet = load_workbook(BytesIO(whole.content)).active
        working = Employee.objects.exclude(left_on__lt=timezone.localdate()).count()
    assert sheet.max_row >= working > 10


def test_the_owner_sets_the_length_in_company_settings(tenant_a):
    client = login(tenant_a, language="en")
    host = {"HTTP_HOST": f"{tenant_a.slug}.localhost"}
    url = reverse("company")
    assert 'name="page_size"' in client.get(url, **host).content.decode()
    with tenant_context(tenant_a.id):
        profile = TenantProfile.objects.get()
    client.post(url, {"display_name": profile.display_name, "accent": profile.accent,
                      "page_size": "50"}, **host)
    with tenant_context(tenant_a.id):
        assert TenantProfile.objects.get().page_size == 50


def test_only_the_choices_are_accepted(tenant_a):
    client = login(tenant_a, language="en")
    host = {"HTTP_HOST": f"{tenant_a.slug}.localhost"}
    with tenant_context(tenant_a.id):
        profile = TenantProfile.objects.get()
    client.post(reverse("company"), {"display_name": profile.display_name, "accent": profile.accent,
                                     "page_size": "100000"}, **host)
    with tenant_context(tenant_a.id):
        assert TenantProfile.objects.get().page_size == 25


def test_a_member_without_the_setting_cannot_change_it(tenant_a, make_member):
    with tenant_context(tenant_a.id):
        assert Membership.objects.exists()
    make_member(tenant_a, "local", "viewer")
    response = login(tenant_a, "local", language="en").post(
        reverse("company"), {"page_size": "100"}, HTTP_HOST=f"{tenant_a.slug}.localhost")
    assert response.status_code in (302, 403)
    with tenant_context(tenant_a.id):
        assert TenantProfile.objects.get().page_size == 25


def test_the_console_has_its_own_length():
    assert "page_size" in PlatformSettingsForm().fields
    assert PlatformSettings._meta.get_field("page_size").default == 25
