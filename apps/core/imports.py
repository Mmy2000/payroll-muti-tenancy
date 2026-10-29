"""Importing rows from Excel, all or nothing (Template Method).

`SheetImport.run()` reads the first sheet, checks the header row, and hands each non-empty row
to `import_row()` inside its own savepoint, collecting every row's problems. If any row fails,
nothing is kept and `ImportFailed` carries the report; the user fixes the file and tries again.

A subclass names its `columns` (the header row holds their keys; other columns are ignored), its
`permission`, and implements `import_row(cells)`, usually by calling the same service function
the screens use. `prepare()` may load lookups once before the rows:

    class EmployeeImport(SheetImport):
        columns = [Column("code", _("Employee code"), required=True, example=101), ...]
        permission = "employees.employee.import"

        def prepare(self):
            self.branches = dict(Branch.objects.values_list("code", "pk"))

        def import_row(self, cells):
            create_employee(EmployeeInput(code=cells.integer("code"), ...), actor=self.actor)
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .errors import AppError


@dataclass(frozen=True)
class Column:
    key: str
    label: object
    required: bool = False
    example: object = ""


@dataclass
class ImportReport:
    created: int = 0
    rows: int = 0
    errors: list[tuple[int, list[str]]] = field(default_factory=list)  # (sheet row, messages)


class ImportFailed(Exception):
    def __init__(self, report: ImportReport):
        super().__init__("import failed")
        self.report = report


class CellError(ValueError):
    """A cell that can't be read: reported as "<column>: <message>" on its row."""

    def __init__(self, key: str, message: str):
        super().__init__(message)
        self.key, self.message = key, message


class Cells:
    """One sheet row by column key, with typed readers. Empty cells give the default."""

    def __init__(self, values: dict):
        self.values = values

    def text(self, key: str) -> str:
        value = self.values.get(key)
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()

    def integer(self, key: str, default=None):
        text = self.text(key)
        if not text:
            return default
        try:
            return int(Decimal(text))
        except (InvalidOperation, ValueError):
            raise CellError(key, _("Enter a whole number.")) from None

    def decimal(self, key: str, default=None):
        text = self.text(key)
        if not text:
            return default
        try:
            return Decimal(text)
        except InvalidOperation:
            raise CellError(key, _("Enter a number.")) from None

    def date(self, key: str):
        value = self.values.get(key)
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value).strip()[:10])
        except ValueError:
            raise CellError(key, _("Enter a date as YYYY-MM-DD.")) from None

    def time(self, key: str):
        value = self.values.get(key)
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.time()
        if isinstance(value, time):
            return value
        text = self.text(key)
        try:
            return time.fromisoformat(text if len(text) > 4 else text.zfill(5))
        except ValueError:
            raise CellError(key, _("Enter a time as HH:MM.")) from None

    def required(self, key: str, value, message: str):
        if value in (None, ""):
            raise CellError(key, message)
        return value


def error_messages(error: AppError) -> list[str]:
    """An AppError's field errors as "field: message" lines (nested fields dotted)."""

    def flatten(prefix, value):
        if isinstance(value, dict):
            for key, inner in value.items():
                yield from flatten(f"{prefix}.{key}" if prefix else key, inner)
        else:
            for message in value if isinstance(value, list) else [value]:
                yield f"{prefix}: {message}" if prefix else str(message)

    return list(flatten("", error.fields)) or [error.message]


class SheetImport:
    columns: list[Column] = []
    permission: str = ""
    sheet_title: str = "data"
    max_rows: int = 2000

    def __init__(self, *, actor=None):
        self.actor = actor

    # --- what a subclass provides ----------------------------------------------------------

    def prepare(self) -> None:
        """Load lookups once, before the rows."""

    def import_row(self, cells: Cells) -> None:
        """Save one row. Raise CellError for a bad cell, or let a service's AppError out."""
        raise NotImplementedError

    # --- the template ----------------------------------------------------------------------

    def template(self) -> bytes:
        book = Workbook()
        sheet = book.active
        sheet.title = self.sheet_title
        sheet.append([c.key for c in self.columns])
        sheet.append([c.example for c in self.columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        guide = book.create_sheet("help")
        guide.append([_("Column"), _("Meaning"), _("Required")])
        for column in self.columns:
            guide.append([column.key, str(column.label), _("yes") if column.required else ""])
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()

    def run(self, upload) -> ImportReport:
        if self.actor is not None and self.permission:
            self.actor.require(self.permission)
        report = ImportReport()
        rows = self._rows(upload, report)
        self.prepare()
        with transaction.atomic():
            for number, values in rows:
                report.rows += 1
                if report.rows > self.max_rows:
                    report.errors.append((number, [_("At most %(count)s rows per file.")
                                                   % {"count": self.max_rows}]))
                    break
                try:
                    with transaction.atomic():
                        self.import_row(Cells(values))
                    report.created += 1
                except CellError as cell:
                    report.errors.append((number, [f"{cell.key}: {cell.message}"]))
                except AppError as error:
                    report.errors.append((number, error_messages(error)))
            if report.errors:
                report.created = 0
                transaction.set_rollback(True)
        if report.errors:
            raise ImportFailed(report)
        return report

    def _rows(self, upload, report: ImportReport):
        """(sheet row number, {key: value}) for every row with something in it."""
        try:
            sheet = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
        except Exception:  # noqa: BLE001 - any unreadable file gets the same answer
            report.errors.append((0, [_("This is not an Excel (.xlsx) file.")]))
            raise ImportFailed(report) from None
        values = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(values, ())]
        missing = [c.key for c in self.columns if c.required and c.key not in header]
        if missing:
            report.errors.append((1, [_("Missing columns: %(names)s")
                                      % {"names": ", ".join(missing)}]))
            raise ImportFailed(report)
        rows = []
        for number, row in enumerate(values, start=2):
            cells = dict(zip(header, row, strict=False))
            if any(v not in (None, "") for v in cells.values()):
                rows.append((number, cells))
        return rows
