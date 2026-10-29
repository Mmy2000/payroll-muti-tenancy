"""Importing employees from Excel: all rows or none."""

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.core.tenancy import tenant_context
from apps.employees import imports
from apps.employees.models import Employee
from conftest import login

pytestmark = pytest.mark.django_db

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet(rows: list[dict]) -> SimpleUploadedFile:
    book = Workbook()
    sheet = book.active
    keys = [c.key for c in imports.EmployeeImport.columns]
    sheet.append(keys)
    for row in rows:
        sheet.append([row.get(key) for key in keys])
    stream = io.BytesIO()
    book.save(stream)
    return SimpleUploadedFile("employees.xlsx", stream.getvalue(), content_type=XLSX)


def _row(code, **extra):
    return {"code": code, "name": f"Person {code}", "branch_code": 1, "hired_on": "2025-01-01",
            "salary": 5000, "starts_at": "08:00", "ends_at": "16:00", "days_off": "fri, sat",
            **extra}


def test_the_template_has_every_column(tenant_a):
    response = login(tenant_a).get("/employees/import/template.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [cell.value for cell in sheet[1]] == [c.key for c in imports.EmployeeImport.columns]


def test_a_good_file_imports_everyone(tenant_a):
    client = login(tenant_a)
    response = client.post("/employees/import/",
                           {"file": _sheet([_row(1), _row(2, days_off="الجمعة")])})
    assert response.status_code == 302
    with tenant_context(tenant_a.id):
        assert Employee.objects.count() == 2
        person = Employee.objects.get(code=1)
        assert person.schedules.get().days_off == [4, 5]
        assert person.pay_history.get().salary == 5000
        assert Employee.objects.get(code=2).schedules.get().days_off == [4]


def test_one_bad_row_keeps_the_whole_file_out(tenant_a):
    client = login(tenant_a, language="en")
    rows = [_row(1), _row(2, branch_code=9), _row(3, starts_at="later"), _row(1)]
    page = client.post("/employees/import/", {"file": _sheet(rows)}).content.decode()
    assert "Nothing was imported" in page
    for number in (3, 4, 5):  # sheet rows; the header is row 1
        assert f'<span class="num">{number}</span>' in page
    with tenant_context(tenant_a.id):
        assert not Employee.objects.exists()


def test_not_an_excel_file(tenant_a):
    junk = SimpleUploadedFile("x.xlsx", b"not a workbook", content_type=XLSX)
    page = login(tenant_a, language="en").post("/employees/import/", {"file": junk})
    assert "not an Excel" in page.content.decode()


def test_importing_needs_its_permission(tenant_a, make_member):
    make_member(tenant_a, "val", "viewer")
    assert login(tenant_a, "val").get("/employees/import/").status_code == 403
