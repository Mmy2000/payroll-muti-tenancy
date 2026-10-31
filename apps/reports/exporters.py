"""Exporting a report's tables (Template Method): `Exporter.export()` walks the tables in one
order (begin, then each table's title, header, rows and totals, then finish), and a subclass
says how each part is written. Reports supply data only; adding a format is one class here and
one line in `EXPORTERS`.

* `PdfExporter`: the A4 sheet (printing/sheet.html) through headless Chromium; sideways when a
  table is wide, with page numbers.
* `XlsxExporter`: an Excel workbook with real numbers and dates (sums work in Excel), money and
  minutes formatted, the sheet right-to-left in Arabic, the header frozen.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from django.utils.translation import get_language_bidi
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .registry import Params, Report, Table

WIDE = 7  # more columns than this: the PDF page goes sideways


def file_stamp(params: Params) -> str:
    return "_".join(v.isoformat() for v in params.chosen.values() if isinstance(v, date))


def landscape(tables: list[Table]) -> bool:
    return max((len(table.columns) for table in tables), default=0) > WIDE


class Exporter:
    extension = ""
    content_type = ""

    def export(self, report: Report, tables: list[Table], params: Params) -> bytes:
        self.begin(report, tables, params)
        for table in tables:
            self.start_table(table)
            for row in table.rows:
                self.row(table, row)
            if table.totals:
                self.totals(table, table.totals)
            self.end_table(table)
        return self.finish()

    def filename(self, report: Report, params: Params) -> str:
        stamp = file_stamp(params)
        return f"{report.code}_{stamp}.{self.extension}" if stamp else \
            f"{report.code}.{self.extension}"

    # --- the steps -------------------------------------------------------------------------

    def begin(self, report: Report, tables: list[Table], params: Params) -> None:
        self.report, self.params = report, params

    def start_table(self, table: Table) -> None:
        pass

    def row(self, table: Table, row: dict) -> None:
        raise NotImplementedError

    def totals(self, table: Table, row: dict) -> None:
        self.row(table, {**row, "_style": "total"})

    def end_table(self, table: Table) -> None:
        pass

    def finish(self) -> bytes:
        raise NotImplementedError


# --- PDF ---------------------------------------------------------------------------------------


def shape(table: Table) -> dict:
    """A table as rows of cells in column order, for the screen and print templates."""
    def cells(row):
        urls = row.get("_urls") or {}
        return [{"value": row.get(c.key), "places": c.places, "kind": c.kind,
                 "url": urls.get(c.key, ""),
                 "negative": c.numeric and (row.get(c.key) or 0) < 0} for c in table.columns]
    return {"title": table.title, "note": table.note, "columns": table.columns,
            "rows": [{"cells": cells(row), "style": row.get("_style", "")} for row in table.rows],
            "totals": cells(table.totals) if table.totals else None}


class PdfExporter(Exporter):
    """The steps only gather the tables; `finish` lays them out on the sheet in one go (the
    page needs all of them to decide on landscape and to repeat headers)."""

    extension, content_type = "pdf", "application/pdf"

    def begin(self, report, tables, params):
        super().begin(report, tables, params)
        self.shaped, self.wide = [], landscape(tables)

    def start_table(self, table):
        self.shaped.append(shape(table))

    def row(self, table, row):
        pass  # already in the shaped table

    def totals(self, table, row):
        pass

    def finish(self) -> bytes:
        from apps.printing.pdf import sheet_pdf

        return sheet_pdf("reports/print.html", {
            "report": self.report, "tables": self.shaped, "params": self.params.chosen,
        }, title=self.report.title, landscape=self.wide)


# --- Excel -------------------------------------------------------------------------------------

FORMATS = {"money": "#,##0.00", "qty": "#,##0.00", "int": "#,##0", "pct": "0.0", "date":
           "yyyy-mm-dd", "minutes": "[h]:mm"}
HEADER_FILL = PatternFill("solid", fgColor="F4F4F5")
BOLD = Font(bold=True)
MINUTES_A_DAY = 24 * 60


def excel_value(value, kind: str):
    """A cell's value as Excel should hold it: numbers as numbers, minutes as a time span."""
    if value is None or value == "":
        return None
    if kind == "minutes":
        return value / MINUTES_A_DAY
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float, date)):
        return value
    return str(value)


class XlsxExporter(Exporter):
    extension = "xlsx"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def begin(self, report, tables, params):
        super().begin(report, tables, params)
        self.book = Workbook()
        self.sheet = self.book.active
        self.sheet.title = str(report.title)[:31]
        self.sheet.sheet_view.rightToLeft = get_language_bidi()
        self.widths: dict[int, int] = {}
        self._write([str(report.title)], bold=True)
        self._write([" → ".join(v.isoformat() for v in params.chosen.values()
                                if isinstance(v, date))])

    def _write(self, values, *, bold=False, fill=None, kinds=None) -> None:
        self.sheet.append(values)
        line = self.sheet.max_row
        for index, value in enumerate(values, start=1):
            cell = self.sheet.cell(row=line, column=index)
            kind = kinds[index - 1] if kinds else "text"
            if kind in FORMATS:
                cell.number_format = FORMATS[kind]
            if bold:
                cell.font = BOLD
            if fill:
                cell.fill = fill
            shown = len(f"{value:,.2f}") if isinstance(value, float) else len(str(value or ""))
            self.widths[index] = max(self.widths.get(index, 8), min(shown + 2, 60))

    def start_table(self, table):
        self.sheet.append([])
        if table.title:
            self._write([str(table.title)], bold=True)
        self._write([str(c.label) for c in table.columns], bold=True, fill=HEADER_FILL)
        if self.sheet.freeze_panes is None:
            self.sheet.freeze_panes = self.sheet.cell(row=self.sheet.max_row + 1, column=1)
        for index, column in enumerate(table.columns, start=1):
            if column.numeric:
                self.sheet.cell(row=self.sheet.max_row, column=index).alignment = \
                    Alignment(horizontal="right")

    def row(self, table, row):
        kinds = [c.kind for c in table.columns]
        self._write([excel_value(row.get(c.key), c.kind) for c in table.columns], kinds=kinds,
                    bold=row.get("_style") in ("subtotal", "total", "heading"))

    def end_table(self, table):
        if table.note:
            self._write([str(table.note)])

    def finish(self) -> bytes:
        for index, width in self.widths.items():
            self.sheet.column_dimensions[get_column_letter(index)].width = width
        stream = io.BytesIO()
        self.book.save(stream)
        return stream.getvalue()


EXPORTERS: dict[str, type[Exporter]] = {"pdf": PdfExporter, "xlsx": XlsxExporter}
